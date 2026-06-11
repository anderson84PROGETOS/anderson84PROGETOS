import tkinter as tk
from tkinter import filedialog, messagebox, Scrollbar
from PIL import Image, ExifTags
import PyPDF2
import docx2txt
import webbrowser
import os
import time
from datetime import datetime

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class MetadataReader:
    def __init__(self, root):
        self.root = root
        self.root.title("🔍 Leitor de Metadados 🔎")
        self.root.geometry("1050x780")
        self.root.wm_state("zoomed")
        self.root.configure(bg="#f0f0f0")

        self.file_path = tk.StringVar()
        self.metadata = {}
        self.geo_data = {}
        self.lat = None
        self.lon = None

        self.create_widgets()

    def create_widgets(self):
        top_frame = tk.Frame(self.root, bg="#f0f0f0")
        top_frame.pack(fill=tk.X, padx=15, pady=12)

        tk.Button(top_frame, text="📁 Selecionar Arquivo", command=self.select_file, 
                  width=20, height=2, bg="#4CAF50", fg="black", font=("Arial", 10, "bold")).pack(side=tk.LEFT)

        tk.Button(top_frame, text="📊 Ler Metadados", command=self.read_metadata, 
                  width=20, height=2, bg="#2196F3", fg="black", font=("Arial", 10, "bold")).pack(side=tk.LEFT, padx=8)

        self.geo_button = tk.Button(top_frame, text="🗺️ Abrir no Google Maps", 
                                   state=tk.DISABLED, command=self.show_map, 
                                   width=26, height=2, bg="#FF5722", fg="black", 
                                   font=("Arial", 10, "bold"))
        self.geo_button.pack(side=tk.LEFT, padx=8)

        tk.Button(top_frame, text="💾 Exportar", command=self.export_metadata, 
                  width=15, height=2, bg="#9C27B0", fg="black", font=("Arial", 10, "bold")).pack(side=tk.RIGHT)

        self.file_label = tk.Label(top_frame, text="Nenhum arquivo selecionado", 
                                  fg="#666666", bg="#f0f0f0", anchor="w", font=("Arial", 10))
        self.file_label.pack(side=tk.LEFT, padx=15, fill=tk.X, expand=True)

        # Área de texto
        text_frame = tk.Frame(self.root, bg="#f0f0f0")
        text_frame.pack(fill=tk.BOTH, expand=True, padx=15, pady=8)

        self.text_area = tk.Text(text_frame, wrap=tk.WORD, font=("Consolas", 10),
                                bg="#ffffff", fg="#1e1e1e", relief=tk.SUNKEN, bd=2)
        scrollbar = tk.Scrollbar(text_frame, orient=tk.VERTICAL, command=self.text_area.yview)
        self.text_area.config(yscrollcommand=scrollbar.set)

        self.text_area.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

    def select_file(self):
        file_path = filedialog.askopenfilename(
            title="Selecione um arquivo",
            filetypes=[("Imagens", "*.jpg *.jpeg *.png *.tiff *.webp"), 
                       ("Todos os arquivos", "*.*")]
        )
        if file_path:
            self.file_path.set(file_path)
            self.file_label.config(text=os.path.basename(file_path))
            self.text_area.delete(1.0, tk.END)
            self.text_area.insert(tk.END, f"✅ Arquivo selecionado: {file_path}\n\n")
            self.geo_button.config(state=tk.DISABLED)

    def read_metadata(self):
        path = self.file_path.get()
        if not path or not os.path.exists(path):
            messagebox.showerror("Erro", "Nenhum arquivo selecionado!")
            return

        self.metadata.clear()
        self.geo_data.clear()
        self.lat = self.lon = None
        self.text_area.delete(1.0, tk.END)

        try:
            _, ext = os.path.splitext(path.lower())
            if ext in ['.jpg', '.jpeg', '.png', '.tiff', '.webp']:
                self.read_image_metadata(path)
            elif ext == '.pdf':
                self.read_pdf_metadata(path)
            elif ext == '.docx':
                self.read_docx_metadata(path)
            else:
                self.read_generic_metadata(path)

            self.display_metadata()
        except Exception as e:
            messagebox.showerror("Erro", f"Erro ao ler metadados:\n{str(e)}")

    def convert_to_decimal(self, gps_val, ref):
        if not gps_val or not ref:
            return None
        try:
            if isinstance(gps_val, (list, tuple)) and len(gps_val) >= 3:
                degrees = float(gps_val[0][0]) / float(gps_val[0][1]) if isinstance(gps_val[0], (list, tuple)) else float(gps_val[0])
                minutes = float(gps_val[1][0]) / float(gps_val[1][1]) if isinstance(gps_val[1], (list, tuple)) else float(gps_val[1])
                seconds = float(gps_val[2][0]) / float(gps_val[2][1]) if isinstance(gps_val[2], (list, tuple)) else float(gps_val[2])

                decimal = degrees + (minutes / 60.0) + (seconds / 3600.0)
                if str(ref).upper() in ['S', 'W']:
                    decimal = -decimal
                return round(decimal, 6)
            return None
        except:
            return None

    def read_image_metadata(self, path):
        try:
            img = Image.open(path)
            self.metadata["Formato"] = img.format or "Desconhecido"
            self.metadata["Dimensões"] = f"{img.size[0]} × {img.size[1]} pixels"
            self.metadata["Modo"] = img.mode
            self.metadata["Caminho"] = path

            exif = img._getexif()
            if exif:
                for tag_id, value in exif.items():
                    tag = ExifTags.TAGS.get(tag_id, f"Tag_{tag_id}")
                    if isinstance(value, bytes):
                        value = value.decode(errors='replace')
                    # Special handling for DateTime
                    if tag in ["DateTime", "DateTimeDigitized", "DateTimeOriginal"]:
                        try:
                            dt = datetime.strptime(value, "%Y:%m:%d %H:%M:%S")
                            value = dt.strftime("%d/%m/%Y %H:%M:%S")
                        except ValueError:
                            pass  # Keep original if parsing fails
                    self.metadata[f"\n\nEXIF: {tag}"] = str(value)

                

                # GPS
                if 34853 in exif:
                    gps_info = exif[34853]
                    for gps_tag_id, gps_value in gps_info.items():
                        gps_tag = ExifTags.GPSTAGS.get(gps_tag_id, f"GPS_Tag_{gps_tag_id}")
                        self.geo_data[gps_tag] = gps_value

                    lat = self.convert_to_decimal(
                        self.geo_data.get('GPSLatitude'),
                        self.geo_data.get('GPSLatitudeRef')
                    )
                    lon = self.convert_to_decimal(
                        self.geo_data.get('GPSLongitude'),
                        self.geo_data.get('GPSLongitudeRef')
                    )

                    if lat is not None and lon is not None:
                        self.lat = lat
                        self.lon = lon
                        self.metadata['\n📍 Coordenadas GPS'] = f"{lat}, {lon}\n"
                        self.geo_button.config(state=tk.NORMAL)
        except Exception as e:
            self.metadata["Erro na leitura de imagem"] = str(e)

    def show_map(self):
        if self.lat is not None and self.lon is not None:
            url = f"https://www.google.com/maps?q={self.lat},{self.lon}&z=18"
            webbrowser.open(url)
        else:
            messagebox.showwarning("Aviso", "Não foi possível obter coordenadas GPS válidas.")

    def read_pdf_metadata(self, path):
        try:
            with open(path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                info = reader.metadata
                if info:
                    for key, value in info.items():
                        clean_key = key.replace('/', '')
                        self.metadata[f"PDF.{clean_key}"] = str(value)
                self.metadata["Número de páginas"] = len(reader.pages)
                self.metadata["Caminho"] = path
        except Exception as e:
            self.metadata["Erro PDF"] = str(e)

    def read_docx_metadata(self, path):
        try:
            text = docx2txt.process(path)
            self.metadata["Número de caracteres"] = len(text)
            self.metadata["Número de palavras"] = len(text.split())
            self.metadata["Caminho"] = path

            if DOCX_AVAILABLE:
                doc = Document(path)
                core = doc.core_properties
                self.metadata["Autor"] = core.author or "Não informado"
                self.metadata["Criado em"] = str(core.created) if core.created else "Não informado"
                self.metadata["Modificado em"] = str(core.modified) if core.modified else "Não informado"
                self.metadata["Título"] = core.title or "Não informado"
        except Exception as e:
            self.metadata["Erro DOCX"] = str(e)

    def read_generic_metadata(self, path):
        try:
            stat = os.stat(path)
            self.metadata["Nome do arquivo"] = os.path.basename(path)
            self.metadata["Caminho completo"] = path
            self.metadata["Tamanho"] = f"{stat.st_size:,} bytes ({stat.st_size / 1024:.2f} KB)"
            self.metadata["Criado em"] = time.ctime(stat.st_ctime)
            self.metadata["Modificado em"] = time.ctime(stat.st_mtime)
            self.metadata["Acessado em"] = time.ctime(stat.st_atime)
        except Exception as e:
            self.metadata["Erro genérico"] = str(e)

    def display_metadata(self):
        self.text_area.insert(tk.END, "📄 METADADOS DO ARQUIVO\n")
        self.text_area.insert(tk.END, "=" * 90 + "\n\n")

        # Formato compacto: Chave: Valor na mesma linha
        for key, value in sorted(self.metadata.items()):
            self.text_area.insert(tk.END, f"{key}: {value}\n")

        if self.geo_data:
            self.text_area.insert(tk.END, "\n📍 Dados GPS Brutos\n==================\n")
            for k, v in sorted(self.geo_data.items()):
                self.text_area.insert(tk.END, f"   {k}: {v}\n")

        self.text_area.insert(tk.END, "\n" + "=" * 90)

    def export_metadata(self):
        if not self.metadata:
            messagebox.showwarning("Aviso", "Nenhum metadado para exportar!")
            return
        save_path = filedialog.asksaveasfilename(defaultextension=".txt")
        if save_path:
            try:
                with open(save_path, 'w', encoding='utf-8') as f:
                    f.write(f"Metadados do arquivo: {self.file_path.get()}\n")
                    f.write("=" * 80 + "\n\n")
                    for key, value in sorted(self.metadata.items()):
                        f.write(f"{key}: {value}\n")
                messagebox.showinfo("Sucesso", f"Exportado para:\n{save_path}")
            except Exception as e:
                messagebox.showerror("Erro", str(e))


def main():
    root = tk.Tk()
    app = MetadataReader(root)
    root.mainloop()

if __name__ == "__main__":
    main()
