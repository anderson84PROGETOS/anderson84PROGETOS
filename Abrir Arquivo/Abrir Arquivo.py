import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext
import json
import csv
import os
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document
from openpyxl import Workbook
import xml.etree.ElementTree as ET

def abrir_arquivo():
    caminho = filedialog.askopenfilename(title="Selecione um arquivo",
                                         filetypes=[("Todos os arquivos", "*.*")])
    if caminho:
        try:
            # Tenta abrir como UTF-8
            with open(caminho, "r", encoding="utf-8") as f:
                conteudo = f.readlines()
        except UnicodeDecodeError:
            try:
                # Se falhar, tenta como Latin-1 (ISO-8859-1)
                with open(caminho, "r", encoding="latin-1") as f:
                    conteudo = f.readlines()
            except Exception as e:
                messagebox.showerror("Erro", f"Não foi possível abrir o arquivo:\n{e}")
                return
        except Exception as e:
            messagebox.showerror("Erro", f"Não foi possível abrir o arquivo:\n{e}")
            return

        entrada.delete("1.0", tk.END)
        entrada.insert(tk.END, "".join(conteudo))
        salvar_botoes(caminho, conteudo)

def salvar_botoes(caminho_original, conteudo):
    def salvar_formato():
        formato = var_formato.get().lower()
        if not formato:
            messagebox.showwarning("Aviso", "Selecione um formato antes de salvar!")
            return

        caminho_salvar = filedialog.asksaveasfilename(
            defaultextension=f".{formato}",
            filetypes=[(f"Arquivos {formato.upper()}", f"*.{formato}")],
            title=f"Salvar como {formato.upper()}"
        )
        if not caminho_salvar:
            return

        try:
            if formato == "json":
                dados = [linha.strip().split(",") for linha in conteudo if linha.strip()]
                headers = dados[0]
                lista_dict = [dict(zip(headers, linha)) for linha in dados[1:]]
                with open(caminho_salvar, "w", encoding="utf-8") as f:
                    json.dump(lista_dict, f, indent=4, ensure_ascii=False)

            elif formato == "csv":
                dados = [linha.strip().split(",") for linha in conteudo if linha.strip()]
                with open(caminho_salvar, "w", newline='', encoding="utf-8") as f:
                    escritor = csv.writer(f)
                    escritor.writerows(dados)

            elif formato == "txt":
                with open(caminho_salvar, "w", encoding="utf-8") as f:
                    f.write("".join(conteudo))

            elif formato == "pdf":
                doc = SimpleDocTemplate(caminho_salvar, pagesize=letter)
                styles = getSampleStyleSheet()
                elementos = [Paragraph(line.strip(), styles["Normal"]) for line in conteudo if line.strip()]
                doc.build(elementos)

            elif formato == "docx":
                doc = Document()
                for linha in conteudo:
                    doc.add_paragraph(linha.strip())
                doc.save(caminho_salvar)

            elif formato == "xlsx":
                wb = Workbook()
                ws = wb.active
                for linha in conteudo:
                    dados = linha.strip().split(",")
                    ws.append(dados)
                wb.save(caminho_salvar)

            elif formato == "html":
                html_content = "<html><body><table border='1'>"
                for linha in conteudo:
                    dados = linha.strip().split(",")
                    html_content += "<tr>" + "".join(f"<td>{dado}</td>" for dado in dados) + "</tr>"
                html_content += "</table></body></html>"
                with open(caminho_salvar, "w", encoding="utf-8") as f:
                    f.write(html_content)

            elif formato == "xml":
                dados = [linha.strip().split(",") for linha in conteudo if linha.strip()]
                headers = dados[0]
                root = ET.Element("root")
                for linha in dados[1:]:
                    item = ET.SubElement(root, "item")
                    for header, valor in zip(headers, linha):
                        campo = ET.SubElement(item, header)
                        campo.text = valor
                tree = ET.ElementTree(root)
                tree.write(caminho_salvar, encoding="utf-8", xml_declaration=True)

            messagebox.showinfo("Sucesso", f"Arquivo salvo em: {caminho_salvar}")
        except Exception as e:
            messagebox.showerror("Erro", f"Não foi possível salvar o arquivo:\n{e}")

    botao_salvar.config(state="normal", command=salvar_formato)

# Interface gráfica
janela = tk.Tk()
janela.title("Abrir Arquivo")
janela.geometry("800x650")

btn_abrir = tk.Button(janela, text="Abrir Arquivo", command=abrir_arquivo,
                      font=("Arial", 12), fg="black", bg="#03fc7f")
btn_abrir.pack(pady=10)

frame_botoes = tk.Frame(janela)
frame_botoes.pack(pady=10)

formatos = ["JSON", "CSV", "TXT", "PDF", "DOCX", "XLSX", "HTML", "XML"]
var_formato = tk.StringVar(janela)
var_formato.set(formatos[0])

menu_formato = tk.OptionMenu(frame_botoes, var_formato, *formatos)
menu_formato.config(font=("Arial", 12), bg="#05ffff", fg="black")
menu_formato["menu"].config(bg="#05ffff", fg="black")
menu_formato.pack(pady=5)

botao_salvar = tk.Button(frame_botoes, text="Salvar", state="disabled",
                         font=("Arial", 12), fg="black", bg="#ff8605")
botao_salvar.pack(pady=5)

# Área de texto com scrollbar
entrada = scrolledtext.ScrolledText(janela, width=120, height=45)
entrada.pack(pady=10)

janela.mainloop()
